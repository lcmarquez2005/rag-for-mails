import json
from src.schemas import MachineRecord, ShiftReport
from src.services.ai import sanitize_json_response
from src.extractor import MockShiftExtractor
from src.exporters import JsonExporter, ExcelExporter
from src.pipeline import IngestionPipeline


def test_machine_record_validations():
    # Caso estándar con llaves principales
    record = MachineRecord(
        machine_id="m102 ",
        downtime_minutes=45,
        approved_parts=1200,
        rejected_parts=15,
        reasons="cambio de molde, defectos estéticos",
        shift_leader="juan pérez",
        shift="Turno 1"
    )

    # Validar normalización y datos extraídos
    assert record.machine_id == "M102"
    assert record.downtime_minutes == 45
    assert record.approved_parts == 1200
    assert record.rejected_parts == 15
    assert record.shift_leader == "Juan Pérez"
    assert record.shift == "Turno 1"
    assert record.reasons == ["cambio de molde", "defectos estéticos"]


def test_sanitize_json_response():
    # Respuesta con markdown envolvente
    markdown_wrapped = """
    Aquí está el resultado:
    ```json
    {
      "records": [
        {
          "machine_id": "M102",
          "downtime_minutes": 45,
          "approved_parts": 1200,
          "rejected_parts": 15,
          "reasons": ["cambio de molde"],
          "shift_leader": "Juan Pérez",
          "shift": "Turno 1"
        }
      ]
    }
    ```
    """
    data = sanitize_json_response(markdown_wrapped)
    assert "records" in data
    assert len(data["records"]) == 1
    assert data["records"][0]["machine_id"] == "M102"


def test_mock_extractor_standard_case():
    text = """Actualización del Turno 1: 
    La máquina M102 estuvo detenida durante 45 minutos debido a un cambio de molde. Total de piezas aprobadas: 1200. Piezas rechazadas: 15 debido a defectos estéticos. El equipo de mantenimiento completó las revisiones preventivas en la M102. 
    Líder de turno: Juan Pérez."""

    extractor = MockShiftExtractor()
    report = extractor.extract(text, source_file="test.txt")

    assert len(report.records) == 1
    rec = report.records[0]
    assert rec.machine_id == "M102"
    assert rec.downtime_minutes == 45
    assert rec.approved_parts == 1200
    assert rec.rejected_parts == 15
    assert rec.shift_leader == "Juan Pérez"


def test_pipeline_e2e_output_json_and_excel(tmp_path):
    json_file = tmp_path / "output.json"
    excel_file = tmp_path / "test_report.xlsx"

    pipeline = IngestionPipeline(force_mock=True)
    pipeline.json_exporter = JsonExporter(output_path=json_file)
    pipeline.excel_exporter = ExcelExporter(output_path=excel_file)

    sample_text = """Actualización del Turno 2:
    Máquina M105: 850 piezas aprobadas, 95 rechazadas por quemadura. Paro de 75 minutos.
    Líder de turno: Maria Rodriguez"""

    result = pipeline.process_text(sample_text, source_file="test_sample.txt")

    assert result.success is True
    assert json_file.exists()
    assert excel_file.exists()
    assert result.report.records[0].machine_id == "M105"

    # Verificar estructura original exacta en output.json
    with open(json_file, "r", encoding="utf-8") as f:
        data = json.load(f)
    assert isinstance(data, list)
    assert len(data) == 1
    record = data[0]
    assert record["machine_id"] == "M105"
    assert record["downtime_minutes"] == 75
    assert record["approved_parts"] == 850
    assert record["rejected_parts"] == 95
    assert record["shift_leader"] == "Maria Rodriguez"
    assert "reasons" in record
    # Sin adaptaciones ni campos ajenos
    assert "sheetId" not in record
    assert "cells" not in record
    assert "scrap_rate_pct" not in record


def test_excel_exporter_idempotency(tmp_path):
    excel_file = tmp_path / "idempotent_report.xlsx"
    exporter = ExcelExporter(output_path=excel_file)

    report = ShiftReport(
        records=[
            MachineRecord(
                machine_id="M102",
                downtime_minutes=45,
                approved_parts=1200,
                rejected_parts=15,
                reasons=["cambio de molde"],
                shift_leader="Juan Pérez",
                shift="Turno 1"
            )
        ],
        raw_source="test raw",
        source_file="gmail_msg_abc123",
        email_id="msg_abc123",
        email_date="2026-09-10 10:00:00"
    )

    # Primera exportación
    exporter.export_report(report)
    import openpyxl
    wb = openpyxl.load_workbook(excel_file)
    ws = wb.active
    assert ws.max_row == 2
    assert ws.max_column == 9
    assert ws.cell(row=1, column=2).value == "Fecha del Correo"
    assert ws.cell(row=1, column=3).value == "Turno"
    assert ws.cell(row=1, column=4).value == "Máquina ID"
    assert ws.cell(row=2, column=1).value == "msg_abc123"
    assert ws.cell(row=2, column=2).value == "2026-09-10 10:00:00"
    assert ws.cell(row=2, column=3).value == "Turno 1"
    assert ws.cell(row=2, column=4).value == "M102"

    # Segunda exportación idéntica (debe omitir duplicado)
    exporter.export_report(report)
    wb2 = openpyxl.load_workbook(excel_file)
    ws2 = wb2.active
    assert ws2.max_row == 2  # No se añade fila duplicada
    assert ws2.max_column == 9


def test_excel_exporter_removes_fecha_proceso_column(tmp_path):
    excel_file = tmp_path / "old_with_fecha_proceso.xlsx"
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([
        "ID de Correo / Origen",
        "Fecha del Correo",
        "Fecha/Hora Proceso",
        "Turno",
        "Máquina ID",
        "Líder de Turno",
        "Paro (Minutos)",
        "Piezas Aprobadas",
        "Piezas Rechazadas",
        "Motivos / Observaciones"
    ])
    ws.append(["msg_old", "2026-09-10 10:00:00", "2026-09-10 10:05:00", "Turno 1", "M102", "Juan", 0, 100, 0, "OK"])
    wb.save(excel_file)

    exporter = ExcelExporter(output_path=excel_file)
    wb_migrated, ws_migrated = exporter._get_or_create_workbook()

    assert ws_migrated.max_column == 9
    headers = [cell.value for cell in ws_migrated[1]]
    assert "Fecha/Hora Proceso" not in headers
    assert headers[1] == "Fecha del Correo"
    assert headers[2] == "Turno"
    assert headers[3] == "Máquina ID"

    row_vals = [cell.value for cell in ws_migrated[2]]
    assert row_vals == ["msg_old", "2026-09-10 10:00:00", "Turno 1", "M102", "Juan", 0, 100, 0, "OK"]


def test_processed_tracker(tmp_path):
    from src.exporters import ProcessedTracker
    tracker_file = tmp_path / "tracker.json"
    tracker = ProcessedTracker(path=tracker_file)

    assert tracker.is_processed("email_123") is False
    tracker.record(email_id="email_123", email_date="2026-09-10", sender="l23@pachuca.tecnm.mx", records_count=2)
    assert tracker.is_processed("email_123") is True
    # Re-registro no duplica
    tracker.record(email_id="email_123")
    with open(tracker_file, "r") as f:
        data = json.load(f)
    assert len(data) == 1


def test_extractor_textual_numbers_and_typos():
    from src.extractor import MockShiftExtractor
    extractor = MockShiftExtractor()
    informal_text = """
    Reporte del turno uno
    Superbisor: Roberto Gomez
    maquna 102
    se sacaron mil doscientas piezas aprovadas y rechasadas quince por rebaba.
    tiempo detenido cuarenta y cinco minutos por ajuste de parametros.
    """
    report = extractor.extract(informal_text)
    assert len(report.records) == 1
    rec = report.records[0]
    assert rec.machine_id == "M102"
    assert rec.shift == "Turno 1"
    assert rec.shift_leader == "Roberto Gomez"
    assert rec.approved_parts == 1200
    assert rec.rejected_parts == 15
    assert rec.downtime_minutes == 45
    assert "rebaba" in rec.reasons
    assert "ajuste de parametros" in rec.reasons


def test_extractor_reasons_and_mil_quinientas():
    from src.extractor import MockShiftExtractor
    extractor = MockShiftExtractor()
    text = "Reporte del turno uno. La maquina M106 estuvo parada cuarenta minutos por falta de materia prima en tolva. Se fabricaron mil quinientas piezas aprobadas y doce piezas rechazadas por espesor irregular. Supervisor a cargo: Sofia Ramirez."
    report = extractor.extract(text)
    assert len(report.records) == 1
    rec = report.records[0]
    assert rec.machine_id == "M106"
    assert rec.shift == "Turno 1"
    assert rec.shift_leader == "Sofia Ramirez"
    assert rec.approved_parts == 1500
    assert rec.rejected_parts == 12
    assert rec.downtime_minutes == 40
    assert "falta de materia prima en tolva" in rec.reasons
    assert "espesor irregular" in rec.reasons

