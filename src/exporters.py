import json
from pathlib import Path
from typing import Dict, Any, Optional, Set, Tuple

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src import config
from src.schemas import ShiftReport, ProcessedEmailLog

# Estilos reutilizables para exportación Excel
_FONT_REGULAR = Font(name="Calibri", size=11)
_THIN_SIDE = Side(style='thin', color='D9D9D9')
_BORDER_THIN = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)
_ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
_ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
_ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
_HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


class ProcessedTracker:
    """
    Registra y consulta el histórico de correos procesados para garantizar
    idempotencia ante múltiples notificaciones o reintentos de Webhook / Pub/Sub.
    """
    def __init__(self, path: Optional[Path] = None):
        target_path = path if path is not None else config.PROCESSED_EMAILS_PATH
        self.path = Path(target_path)
        self.path.parent.mkdir(parents=True, exist_ok=True)

    def _load_data(self) -> list:
        if not self.path.exists():
            return []
        try:
            with open(self.path, "r", encoding="utf-8") as f:
                data = json.load(f)
                return data if isinstance(data, list) else []
        except Exception:
            return []

    def is_processed(self, email_id: str) -> bool:
        if not email_id:
            return False
        return any(item.get("email_id") == email_id for item in self._load_data())

    def record(
        self,
        email_id: str,
        email_date: Optional[str] = None,
        sender: Optional[str] = None,
        subject: Optional[str] = None,
        records_count: int = 0
    ) -> None:
        if not email_id:
            return
        data = self._load_data()
        if any(item.get("email_id") == email_id for item in data):
            return
        log = ProcessedEmailLog(
            email_id=email_id,
            email_date=email_date,
            sender=sender,
            subject=subject,
            records_count=records_count
        )
        data.append(log.model_dump(mode="json"))
        with open(self.path, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False)


class ExcelExporter:
    """
    Gestiona la exportación de los datos limpios a un libro Excel (.xlsx)
    con verificación de idempotencia y trazabilidad por ID y Fecha de correo.
    """
    HEADERS = [
        "ID de Correo / Origen",
        "Fecha del Correo",
        "Turno",
        "Máquina ID",
        "Líder de Turno",
        "Paro (Minutos)",
        "Piezas Aprobadas",
        "Piezas Rechazadas",
        "Motivos / Observaciones"
    ]

    def __init__(self, output_path: Optional[Path] = None):
        target_path = output_path if output_path is not None else config.EXCEL_OUTPUT_PATH
        self.output_path = Path(target_path)
        self.output_path.parent.mkdir(parents=True, exist_ok=True)

    def export_report(self, report: ShiftReport) -> Path:
        wb, ws = self._get_or_create_workbook()

        email_id_str = report.email_id or report.source_file or "Entrada Directa"
        email_date_str = report.email_date or "N/A"

        # Trazabilidad anti-duplicados: obtener llaves ya existentes en la hoja
        existing_keys = self._get_existing_keys(ws)
        rows_added = 0

        for record in report.records:
            record_key = (str(email_id_str).strip(), str(record.machine_id).strip())
            if record_key in existing_keys:
                continue

            existing_keys.add(record_key)

            row_data = [
                email_id_str,
                email_date_str,
                record.shift or "N/A",
                record.machine_id,
                record.shift_leader,
                record.downtime_minutes,
                record.approved_parts,
                record.rejected_parts,
                ", ".join(record.reasons) if record.reasons else "N/A"
            ]

            ws.append(row_data)
            row_idx = ws.max_row
            rows_added += 1

            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = _BORDER_THIN
                cell.font = _FONT_REGULAR
                if col_idx in (1, 2, 3, 4):
                    cell.alignment = _ALIGN_CENTER
                elif col_idx in (6, 7, 8):
                    cell.alignment = _ALIGN_RIGHT
                else:
                    cell.alignment = _ALIGN_LEFT

        if rows_added > 0:
            self._auto_adjust_columns(ws)
            wb.save(self.output_path)

        return self.output_path

    def _get_existing_keys(self, ws) -> Set[Tuple[str, str]]:
        """Recupera los pares (ID Correo, Máquina ID) existentes para evitar duplicación."""
        keys = set()
        header_vals = [str(cell.value or '').strip() for cell in ws[1]]
        machine_col_idx = 3  # default: columna 4 (0-indexed 3)
        if "Máquina ID" in header_vals:
            machine_col_idx = header_vals.index("Máquina ID")

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) <= machine_col_idx or row[0] is None:
                continue
            id_val = str(row[0]).strip()
            machine_val = str(row[machine_col_idx]).strip()
            keys.add((id_val, machine_val))
        return keys

    def _get_or_create_workbook(self):
        if self.output_path.exists():
            try:
                wb = openpyxl.load_workbook(self.output_path)
                ws = wb.active
                header_vals = [str(cell.value or '').strip() for cell in ws[1]]
                if "Fecha/Hora Proceso" in header_vals:
                    wb, ws = self._remove_fecha_proceso_column(wb, ws)
                return wb, ws
            except Exception:
                pass

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reporte de Moldeo"

        ws.append(self.HEADERS)
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.row_dimensions[1].height = 26
        for col_idx in range(1, len(self.HEADERS) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        wb.save(self.output_path)
        return wb, ws

    def _remove_fecha_proceso_column(self, wb, ws):
        """Elimina la columna 'Fecha/Hora Proceso' preservando los datos existentes."""
        header_vals = [str(cell.value or '').strip() for cell in ws[1]]
        if "Fecha/Hora Proceso" in header_vals:
            col_idx = header_vals.index("Fecha/Hora Proceso") + 1  # 1-indexed
            ws.delete_cols(col_idx)
            self._auto_adjust_columns(ws)
            wb.save(self.output_path)
        return wb, ws

    def _auto_adjust_columns(self, ws):
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


class JsonExporter:
    """
    Exporta los registros procesados en formato JSON limpio con la estructura
    original definida, con prevención activa de registros duplicados.
    """
    def __init__(self, output_path: Optional[Path] = None):
        target_path = output_path if output_path is not None else config.OUTPUT_JSON_PATH
        self.output_path = Path(target_path)
        self.output_path.parent.mkdir(parents=True, exist_ok=True)

    @staticmethod
    def _record_sig(r) -> tuple:
        if isinstance(r, dict):
            return (r.get("machine_id"), r.get("shift_leader"), r.get("downtime_minutes"), r.get("approved_parts"), r.get("rejected_parts"))
        return (r.machine_id, r.shift_leader, r.downtime_minutes, r.approved_parts, r.rejected_parts)

    def export_report(self, report: ShiftReport) -> Dict[str, Any]:
        """Guarda o anexa registros omitiendo duplicados exactos."""
        existing_data = []
        if self.output_path.exists():
            try:
                with open(self.output_path, "r", encoding="utf-8") as f:
                    content = json.load(f)
                    existing_data = content if isinstance(content, list) else ([content] if isinstance(content, dict) else [])
            except Exception:
                existing_data = []

        existing_signatures = {self._record_sig(item) for item in existing_data}
        new_items = []
        for rec in report.records:
            rec_sig = self._record_sig(rec)
            if rec_sig in existing_signatures:
                continue

            existing_signatures.add(rec_sig)
            new_items.append({
                "machine_id": rec.machine_id,
                "downtime_minutes": rec.downtime_minutes,
                "approved_parts": rec.approved_parts,
                "rejected_parts": rec.rejected_parts,
                "reasons": rec.reasons,
                "shift_leader": rec.shift_leader
            })

        existing_data.extend(new_items)

        with open(self.output_path, "w", encoding="utf-8") as f:
            json.dump(existing_data, f, indent=2, ensure_ascii=False)

        return {
            "status": "success",
            "output_file": str(self.output_path),
            "records_inserted": len(new_items),
            "payload_sample": new_items
        }


class GoogleSheetsExporter:
    """
    Gestiona la sincronización en tiempo real con una hoja única de Google Sheets
    ubicada en la carpeta designada de Google Drive, con deduplicación idempotente
    y trazabilidad por ID de Correo y Máquina.
    """
    HEADERS = [
        "ID de Correo / Origen",
        "Fecha del Correo",
        "Turno",
        "Máquina ID",
        "Líder de Turno",
        "Paro (Minutos)",
        "Piezas Aprobadas",
        "Piezas Rechazadas",
        "Motivos / Observaciones"
    ]

    def __init__(
        self,
        sheet_id: Optional[str] = None,
        tab_name: Optional[str] = None,
        client: Optional[Any] = None,
    ):
        self.sheet_id = (sheet_id if sheet_id is not None else config.GOOGLE_SHEET_ID).strip()
        self.tab_name = (tab_name if tab_name is not None else config.GOOGLE_SHEET_TAB_NAME).strip() or "Reportes"
        self.client = client

    def _get_client(self):
        if self.client is None:
            from src.sheets_client import GoogleSheetsClient
            self.client = GoogleSheetsClient()
        return self.client

    def export_report(self, report: ShiftReport) -> Optional[str]:
        """
        Exporta el reporte a Google Sheets de manera idempotente.
        Retorna la URL web de edición del archivo Google Sheet o None si no está configurado.
        """
        if not self.sheet_id:
            return None

        client = self._get_client()
        email_id_str = report.email_id or report.source_file or "Entrada Directa"
        email_date_str = report.email_date or "N/A"

        try:
            existing_keys = client.get_existing_keys(self.sheet_id, self.tab_name)
            rows_to_insert = []

            for record in report.records:
                record_key = (str(email_id_str).strip(), str(record.machine_id).strip())
                if record_key in existing_keys:
                    continue

                existing_keys.add(record_key)
                rows_to_insert.append([
                    email_id_str,
                    email_date_str,
                    record.shift or "N/A",
                    record.machine_id,
                    record.shift_leader,
                    record.downtime_minutes,
                    record.approved_parts,
                    record.rejected_parts,
                    ", ".join(record.reasons) if record.reasons else "N/A"
                ])

            if rows_to_insert:
                client.append_rows(self.sheet_id, rows_to_insert, tab_name=self.tab_name)

            return f"https://docs.google.com/spreadsheets/d/{self.sheet_id}/edit"
        except Exception as e:
            import logging
            logger = logging.getLogger("rag_mails.exporters.sheets")
            logger.error(f"Error sincronizando con Google Sheets: {e}", exc_info=True)
            return None

